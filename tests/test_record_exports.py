"""Observation passes the step, approval signatures (with images) under a record, PDF / Excel of one record and when
those buttons show (always / only when approved; page default in code, overridden by the approval type)."""
import io

from django.core.files.base import ContentFile
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image

from core import approvals
from core.approvals.signatures import approval_signatures
from core.crud.registry import get_resource
from core.models import Approval, ApprovalStep, ApprovalType

from tests.dummyapp.models import Garment, GarmentLine

from .base import CoreTestCase, make_user
from .test_custom_reports import pdf_text

MODAL = {"HTTP_HX_REQUEST": "true", "HTTP_HX_TARGET": "core-modal-content"}
VIEW = ["dummyapp.view_garment"]


def png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (120, 40), "navy").save(buffer, "PNG")
    return buffer.getvalue()


class ApprovalFixture(CoreTestCase):
    def setUp(self):
        super().setUp()
        approvals.sync_all_chains()
        self.maker = make_user("maker@example.com", self.company_a, perms=[*VIEW, "dummyapp.add_garment", "dummyapp.change_garment"])
        self.merch = make_user("merch@example.com", self.company_a, perms=[*VIEW, "dummyapp.change_garment"])
        self.boss = make_user("boss@example.com", self.company_a, perms=[*VIEW, "dummyapp.approve_garment"])
        self.g = Garment.objects.create(company=self.company_a, name="Tee", style_no="T1", price=4)
        GarmentLine.objects.create(garment=self.g, description="Front panel", qty=7)
        self.atype = ApprovalType.objects.get(code="garment")
        ApprovalStep.objects.update(allow_observation=True)

    def approve_all(self):
        req = approvals.submit_for_approval(self.maker, self.g)
        approvals.submit_approval_step(self.merch, req.pk, "approve")
        approvals.submit_approval_step(self.boss, req.pk, "approve")
        return req


class ObservationPassesTheStepTests(ApprovalFixture):
    def test_observation_moves_the_request_to_the_next_step(self):
        req = approvals.submit_for_approval(self.maker, self.g)
        first = req.current_sequence
        approvals.submit_approval_step(self.merch, req.pk, "observe", comment="Check the FOB")
        req.refresh_from_db()
        self.assertEqual(req.status, "pending")
        self.assertGreater(req.current_sequence, first)
        self.assertTrue(req.history.filter(action=Approval.OBSERVE, sequence=first).exists())
        self.assertTrue(approvals.can_act(self.boss, req))
        self.assertFalse(approvals.can_act(self.merch, req))

    def test_observation_on_the_final_step_approves(self):
        req = approvals.submit_for_approval(self.maker, self.g)
        approvals.submit_approval_step(self.merch, req.pk, "approve")
        approvals.submit_approval_step(self.boss, req.pk, "observe", comment="Fine, mind the shade")
        req.refresh_from_db()
        self.g.refresh_from_db()
        self.assertEqual(req.status, "approved")
        self.assertTrue(self.g.is_approval)

    def test_observed_step_counts_as_signed(self):
        req = approvals.submit_for_approval(self.maker, self.g)
        approvals.submit_approval_step(self.merch, req.pk, "observe", comment="ok")
        sigs = approval_signatures(self.g)
        self.assertEqual([s.name for s in sigs], ["maker@example.com", "merch@example.com", ""])

    def test_queue_post_reports_the_move(self):
        req = approvals.submit_for_approval(self.maker, self.g)
        client = self.login(self.merch)
        r = client.post(reverse("core:approval_detail", args=[req.pk]),
                        {"action": "observe", "comment": "note", "expected_sequence": req.current_sequence}, follow=True)
        self.assertContains(r, "the request moved on")


class SignatureRowTests(ApprovalFixture):
    def test_signature_label_and_image_under_the_detail(self):
        step = ApprovalStep.objects.get(code="merch")
        step.signature_label = "Checked By"
        step.save()
        self.merch.profile.user_sign.save("merch.png", ContentFile(png()))
        self.approve_all()
        html = self.login(self.maker).get(reverse("core:resource_detail", args=["garments", self.g.pk]), **MODAL).content.decode()
        for text in ("Prepared By", "Checked By", "Manager", "merch@example.com", "boss@example.com", "data:image/png;base64,"):
            self.assertIn(text, html)
        sigs = approval_signatures(self.g)
        self.assertEqual(sigs[1].label, "Checked By")
        self.assertTrue(sigs[1].image)
        self.assertIsNone(sigs[2].image)

    def test_pending_steps_show_as_pending(self):
        approvals.submit_for_approval(self.maker, self.g)
        html = self.login(self.maker).get(reverse("core:resource_detail", args=["garments", self.g.pk]), **MODAL).content.decode()
        self.assertIn("core-signatures", html)
        self.assertIn("Pending", html)

    def test_no_row_before_submission_or_after_withdrawal(self):
        client = self.login(self.maker)
        url = reverse("core:resource_detail", args=["garments", self.g.pk])
        self.assertNotIn("core-signatures", client.get(url, **MODAL).content.decode())
        approvals.submit_for_approval(self.maker, self.g)
        approvals.cancel(self.maker, self.g)
        self.assertNotIn("core-signatures", client.get(url, **MODAL).content.decode())

    def test_full_detail_page_has_the_row_too(self):
        self.approve_all()
        html = self.login(self.maker).get(reverse("core:resource_detail", args=["garments", self.g.pk])).content.decode()
        self.assertIn("core-signatures", html)
        self.assertIn(reverse("core:resource_record_export", args=["garments", self.g.pk, "pdf"]), html)


class RecordExportTests(ApprovalFixture):
    def url(self, ext):
        return reverse("core:resource_record_export", args=["garments", self.g.pk, ext])

    def detail(self, user):
        return self.login(user).get(reverse("core:resource_detail", args=["garments", self.g.pk]), **MODAL).content.decode()

    def test_buttons_show_all_the_time_by_default(self):
        html = self.detail(self.maker)
        self.assertIn(self.url("pdf"), html)
        self.assertIn(self.url("xlsx"), html)

    def test_pdf_has_fields_lines_and_signatures(self):
        self.merch.profile.user_sign.save("merch.png", ContentFile(png()))
        self.approve_all()
        r = self.login(self.maker).get(self.url("pdf"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertIn(f'garments-{self.g.pk}.pdf', r["Content-Disposition"])
        text = pdf_text(r.content)
        for expected in ("Garment - Tee", "T1", "Front panel", "Prepared By", "Manager", "boss@example.com"):
            self.assertIn(expected, text)

    def test_excel_has_fields_lines_and_approvals(self):
        self.merch.profile.user_sign.save("merch.png", ContentFile(png()))
        self.approve_all()
        r = self.login(self.maker).get(self.url("xlsx"))
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        values = [c for row in ws.iter_rows(values_only=True) for c in row if c not in (None, "")]
        for expected in ("Garment - Tee", "T1", "Front panel", 7, "Approvals", "Prepared By", "merch@example.com"):
            self.assertIn(expected, values)
        self.assertEqual(len(ws._images), 1)                    # the merchandiser's signature

    def test_only_when_approved(self):
        self.atype.export_when = ApprovalType.WHEN_APPROVED
        self.atype.save()
        client = self.login(self.maker)
        self.assertNotIn(self.url("pdf"), self.detail(self.maker))
        self.assertEqual(client.get(self.url("pdf")).status_code, 404)
        req = approvals.submit_for_approval(self.maker, self.g)
        approvals.submit_approval_step(self.merch, req.pk, "approve")
        self.assertEqual(client.get(self.url("xlsx")).status_code, 404)       # still pending
        approvals.submit_approval_step(self.boss, req.pk, "approve")
        self.assertIn(self.url("pdf"), self.detail(self.maker))
        self.assertEqual(client.get(self.url("pdf")).status_code, 200)

    def test_approval_type_picks_the_buttons(self):
        self.atype.export_buttons = ApprovalType.EXPORT_PDF
        self.atype.save()
        client = self.login(self.maker)
        self.assertEqual(client.get(self.url("pdf")).status_code, 200)
        self.assertEqual(client.get(self.url("xlsx")).status_code, 404)
        self.atype.export_buttons = ApprovalType.EXPORT_NONE
        self.atype.save()
        self.assertNotIn("resource_record_export", self.detail(self.maker))
        self.assertEqual(client.get(self.url("pdf")).status_code, 404)

    def test_code_default_applies_when_the_type_says_page_default(self):
        cfg = get_resource("garments")
        old = cfg.detail_exports, cfg.detail_export_when
        try:
            cfg.detail_exports, cfg.detail_export_when = ("xlsx",), "approved"
            client = self.login(self.maker)
            self.assertEqual(client.get(self.url("xlsx")).status_code, 404)
            self.approve_all()
            self.assertEqual(client.get(self.url("xlsx")).status_code, 200)
            self.assertEqual(client.get(self.url("pdf")).status_code, 404)
        finally:
            cfg.detail_exports, cfg.detail_export_when = old

    def test_records_without_approval_ignore_the_approved_rule(self):
        from tests.dummyapp.models import Fabric

        cfg = get_resource("fabrics")
        old = cfg.detail_export_when
        fabric = Fabric.objects.create(company=self.company_a, name="Jersey")
        user = make_user("fab@example.com", self.company_a, perms=["dummyapp.view_fabric"])
        try:
            cfg.detail_export_when = "approved"
            r = self.login(user).get(reverse("core:resource_record_export", args=["fabrics", fabric.pk, "pdf"]))
            self.assertEqual(r.status_code, 200)
        finally:
            cfg.detail_export_when = old

    def test_needs_access_to_the_record(self):
        outsider = make_user("out@example.com", self.company_b, perms=VIEW)
        self.assertEqual(self.login(outsider).get(self.url("pdf")).status_code, 404)
        no_perm = make_user("noperm@example.com", self.company_a)
        self.assertIn(self.login(no_perm).get(self.url("pdf")).status_code, (403, 404))

    def test_unknown_format_is_404(self):
        self.assertEqual(self.login(self.maker).get(self.url("csv")).status_code, 404)

    def test_export_is_audited(self):
        from core.models import AuditLog

        self.login(self.maker).get(self.url("pdf"))
        self.assertTrue(AuditLog.objects.filter(message__startswith="export garments record pdf").exists())
