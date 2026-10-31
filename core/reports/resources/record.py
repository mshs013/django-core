"""One record as a document: what its detail page shows (fields, inline tables, approval signatures) as a PDF or an Excel workbook,
for the PDF / Excel buttons of the detail page.

Which buttons a record offers: ``register_resource(..., detail_exports=("pdf", "xlsx"), detail_export_when="always")`` in code,
overridden per approval type (Approval types > *Export buttons* / *Show export buttons*). ``"approved"`` hides the buttons -- and
refuses the download -- until the record's approval is final; it only applies to records that have an approval.
"""
from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal
from xml.sax.saxutils import escape

from django.utils.html import strip_tags

from core.reports.letterhead import format_when

from .formats import available_formats
from .pdf import build_branded

HEAD_COLOR = "#206bc4"


def detail_export_formats(cfg, obj) -> tuple[str, ...]:
    """The export formats the detail page of ``obj`` offers right now (empty: no buttons, downloads refused)."""
    from core.approvals.errors import ApprovalError
    from core.approvals.registry import target_for
    from core.approvals.service import status_for
    from core.models import ApprovalType

    if not cfg.allow_export:
        return ()
    formats, when = cfg.detail_exports, cfg.detail_export_when
    try:
        target = target_for(obj)
    except ApprovalError:
        target = None
    if target is not None:
        atype = ApprovalType.objects.filter(code=target.chain).first()
        if atype is not None:
            chosen = atype.export_formats()
            formats = formats if chosen is None else chosen
            when = atype.export_when or when
        if when == "approved" and status_for(obj) != "approved":
            return ()
    return tuple(f for f in formats if f == "pdf" or f in available_formats())


def plain(value) -> str:
    """Display text without markup (detail values may be badges or links)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return strip_tags(str(value)).strip()


def _cell_value(value):
    """Numbers stay numbers in Excel; everything else becomes its display text."""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return value
    return plain(value)


# ------------------------------------------------------------------------------------------------------------------- PDF
def record_pdf(title: str, details, children, signatures, meta: dict) -> bytes:
    """Fields in a three-column grid, each inline as a table, the signature row at the bottom of the last page."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("rec_cell", fontSize=8, leading=10)
    head = small.clone("rec_head", textColor=colors.white)
    heading = styles["BodyText"].clone("rec_heading", fontSize=9.5, leading=12, spaceBefore=4)
    body = []
    cells = [Paragraph(f"<b>{escape(str(label))} :</b> {escape(plain(value)) or '-'}", small) for label, value, *_rest in details]
    if cells:
        per_row = 3
        cells += [""] * (-len(cells) % per_row)
        grid = Table([cells[i:i + per_row] for i in range(0, len(cells), per_row)], hAlign="LEFT")
        grid.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 2),
                                  ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        body += [grid, Spacer(1, 8)]
    for child in children:
        label = getattr(child["inline"], "label_plural", "") or child["inline"].label
        body.append(Paragraph(f"<b>{escape(str(label))}</b>", heading))
        data = [[Paragraph(f"<b>{escape(str(h))}</b>", head) for h in ["#", *child["headers"]]]]
        data += [[Paragraph(escape(plain(v)), small) for v in [n, *row]] for n, row in enumerate(child["items"], 1)]
        if not child["items"]:
            data.append([Paragraph("None.", small)] + [""] * len(child["headers"]))
        table = Table(data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(HEAD_COLOR)),
                                   ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9aa0ac")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        body += [table, Spacer(1, 8)]
    wide = any(len(c["headers"]) > 6 for c in children)
    meta = {**meta, "signatures": list(signatures), "pagesize": landscape(A4) if wide else A4}
    buffer = io.BytesIO()
    build_branded(buffer, title, body, meta)
    return buffer.getvalue()


# ------------------------------------------------------------------------------------------------------------------- Excel
def record_xlsx(title: str, details, children, signatures) -> bytes:
    """One sheet: title, the fields as label / value rows, each inline as a table, then the approvals (with signature images)."""
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "".join(c for c in title if c not in "[]:*?/\\")[:31] or "Record"
    bold, section = Font(bold=True), Font(bold=True, size=12)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=HEAD_COLOR.lstrip("#").upper())
    line = Side(style="thin", color="9AA0AC")
    border = Border(left=line, right=line, top=line, bottom=line)
    widths: dict[int, int] = {}

    def put(row, col, value, font=None, framed=False, header=False):
        cell = ws.cell(row=row, column=col, value=_cell_value(value))
        if font is not None:
            cell.font = font
        if header:
            cell.font, cell.fill = head_font, head_fill
        if framed or header:
            cell.border = border
        if isinstance(cell.value, dt.datetime):
            cell.number_format = "yyyy-mm-dd hh:mm"
        widths[col] = max(widths.get(col, 8), min(len(str(cell.value or "")) + 2, 60))
        return cell

    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    row = 3
    for label, value, *_rest in details:
        put(row, 1, str(label), bold)
        put(row, 2, value)
        row += 1
    for child in children:
        row += 1
        inline = child["inline"]
        ws.cell(row=row, column=1, value=str(getattr(inline, "label_plural", "") or inline.label)).font = section
        row += 1
        for col, header in enumerate(["#", *child["headers"]], 1):
            put(row, col, str(header), header=True)
        row += 1
        for n, items in enumerate(child["items"], 1):
            for col, value in enumerate([n, *items], 1):
                put(row, col, value, framed=True)
            row += 1
    if signatures:
        row += 1
        ws.cell(row=row, column=1, value="Approvals").font = section
        row += 1
        for col, header in enumerate(("Step", "Name", "Date", "Signature"), 1):
            put(row, col, header, header=True)
        row += 1
        for sig in signatures:
            put(row, 1, sig.label, framed=True)
            put(row, 2, sig.name, framed=True)
            put(row, 3, format_when(sig.when), framed=True)
            put(row, 4, "", framed=True)
            if _add_image(ws, sig.image, f"D{row}"):
                ws.row_dimensions[row].height = 36
                widths[4] = max(widths.get(4, 8), 22)
            row += 1
    from openpyxl.utils import get_column_letter

    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _add_image(ws, data: bytes | None, anchor: str) -> bool:
    if not data:
        return False
    try:
        from openpyxl.drawing.image import Image

        image = Image(io.BytesIO(data))
        scale = min(150 / float(image.width or 1), 44 / float(image.height or 1))
        image.width, image.height = image.width * scale, image.height * scale
        ws.add_image(image, anchor)
        return True
    except Exception:  # noqa: BLE001 - a broken signature image must not break the workbook
        return False


def record_bytes(cfg, obj, ext: str, details, children, signatures, meta: dict) -> tuple[bytes, str, str]:
    """(data, file name, content type) of one record in ``ext`` ("pdf" | "xlsx")."""
    title = f"{cfg.label} - {obj}"
    name = f"{cfg.slug}-{obj.pk}.{ext}"
    if ext == "pdf":
        return record_pdf(title, details, children, signatures, meta), name, "application/pdf"
    return (record_xlsx(title, details, children, signatures), name,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
